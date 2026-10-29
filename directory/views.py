from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.conf import settings
from django.core.management.base import CommandError
from django.db import IntegrityError, models, transaction
from django.http import Http404, FileResponse, HttpResponse, HttpResponseBadRequest, HttpResponseForbidden, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.http import require_GET, require_http_methods, require_POST
from io import BytesIO
from openpyxl import load_workbook
from pathlib import Path

from .management.commands.import_executors import headers_are_supported, read_xlsx_safe, save_prepared, validate_rows

from .forms import ExecutorForm, ProjectForm
from .models import (
    Executor,
    ExecutorComment,
    ExecutorSpecialty,
    ExecutorStatus,
    ObjectType,
    Project,
    ProjectStatus,
    ProjectSelection,
    ProjectSpecialtyNeed,
    SelectionNegotiation,
    SelectionAttachment,
    SelectionStatus,
    Software,
    Specialty,
)

@login_required
def home(request):
    context = {
        "executors_count": Executor.objects.count(),
        "projects_count": Project.objects.count(),
        "selections_count": ProjectSelection.objects.count(),
        "specialties_count": Specialty.objects.count(),
    }

    return render(request, "directory/home.html", context)

def format_money(value):
    if value is None:
        return None
    return f"{value:,}".replace(",", " ")

def get_selection_status_css_class(status_name):
    normalized_name = (status_name or "").strip().lower()

    if normalized_name == "новый":
        return "selection-status-new"

    if normalized_name == "рассматривает":
        return "selection-status-review"

    if normalized_name in {"готов", "утвержден"}:
        return "selection-status-ready"

    if normalized_name == "отказ":
        return "selection-status-rejected"

    return "selection-status-default"

def get_need_status(selection_list):
    has_ready = any(
        selection.status and selection.status.name == "Утвержден"
        for selection in selection_list
    )

    if has_ready:
        return {
            "label": "Закрыт",
            "css_class": "need-status-closed",
        }

    if selection_list:
        return {
            "label": "В работе",
            "css_class": "need-status-in-progress",
        }

    return {
        "label": "Нет кандидатов",
        "css_class": "need-status-empty",
    }

@login_required
def executor_search(request):
    return redirect("executor_list")


@login_required
def executor_import_template(request):
    path = Path(settings.BASE_DIR) / "directory" / "static" / "directory" / "files" / "executor_import_template.xlsx"
    return FileResponse(
        path.open("rb"),
        as_attachment=True,
        filename="Шаблон_импорта_исполнителей.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@login_required
@require_POST
def export_executors(request):
    if not (request.user.is_staff or request.user.is_superuser):
        return HttpResponseForbidden("Выгрузка доступна только администратору.")

    executor_ids = []
    for value in request.POST.getlist("executor_ids"):
        try:
            executor_id = int(value)
        except (TypeError, ValueError):
            continue
        if executor_id not in executor_ids:
            executor_ids.append(executor_id)

    if not executor_ids:
        return HttpResponseBadRequest("Выберите хотя бы одного исполнителя.")

    executors_by_id = {
        executor.id: executor
        for executor in Executor.objects.filter(id__in=executor_ids)
        .select_related("employment_type", "status", "contact_source")
        .prefetch_related("software", "executor_specialties__specialty")
    }
    executors = [executors_by_id[executor_id] for executor_id in executor_ids if executor_id in executors_by_id]
    if not executors:
        return HttpResponseBadRequest("Выбранные исполнители не найдены.")

    template_path = Path(settings.BASE_DIR) / "directory" / "static" / "directory" / "files" / "executor_import_template.xlsx"
    workbook = load_workbook(template_path)
    worksheet = workbook["Для импорта"]

    for row_number, executor in enumerate(executors, 2):
        worksheet.cell(row_number, 1, executor.last_name)
        worksheet.cell(row_number, 2, executor.first_name)
        worksheet.cell(row_number, 3, executor.middle_name)
        worksheet.cell(row_number, 4, executor.phone)
        worksheet.cell(row_number, 5, executor.email)
        worksheet.cell(row_number, 6, str(executor.employment_type or ""))
        worksheet.cell(row_number, 7, str(executor.status))
        worksheet.cell(row_number, 8, str(executor.contact_source or ""))
        worksheet.cell(row_number, 9, executor.source_comment)
        worksheet.cell(row_number, 10, executor.city)
        worksheet.cell(row_number, 11, executor.work_start_year)
        worksheet.cell(row_number, 12, executor.birth_year)
        worksheet.cell(
            row_number,
            13,
            "; ".join(sorted(program.name for program in executor.software.all())),
        )
        worksheet.cell(row_number, 14, executor.software_comment)
        worksheet.cell(
            row_number,
            15,
            "; ".join(sorted(item.specialty.code for item in executor.executor_specialties.all())),
        )
        worksheet.cell(row_number, 16, executor.general_comment)

    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="executors_export.xlsx"'
    return response


@login_required
def executor_import(request):
    context = {}
    if request.method == "POST":
        upload = request.FILES.get("file")
        if not upload:
            context["error"] = "Выберите Excel-файл для проверки."
        elif not upload.name.lower().endswith(".xlsx"):
            context["error"] = "Поддерживаются только файлы XLSX."
        elif upload.size > 10 * 1024 * 1024:
            context["error"] = "Размер файла не должен превышать 10 МБ."
        else:
            try:
                fieldnames, rows = read_xlsx_safe(upload)
                if not headers_are_supported(fieldnames):
                    raise CommandError("Неверный набор или порядок столбцов.")
                if not rows:
                    raise CommandError("Файл не содержит исполнителей.")
                prepared = validate_rows(rows)
                request.session["executor_import_rows"] = rows
                context["preview_rows"] = [
                    {
                        "name": " ".join(filter(None, [row["Фамилия"], row["Имя"], row["Отчество"]])),
                        "phone": row["Телефон"],
                        "email": row["Почта"],
                        "status": row["Статус"],
                        "source": row["Источник контакта"],
                        "software": row["ПО"],
                        "specialties": row["Специальности"],
                    }
                    for row in rows
                ]
                context["preview_count"] = len(prepared)
            except CommandError as exc:
                request.session.pop("executor_import_rows", None)
                context["error"] = str(exc)
    return render(request, "directory/executor_import.html", context)


@login_required
@require_POST
def executor_import_confirm(request):
    rows = request.session.get("executor_import_rows")
    if not rows:
        messages.error(request, "Нет проверенного файла для импорта. Загрузите файл повторно.")
        return redirect("executor_import")
    try:
        with transaction.atomic():
            list(ExecutorStatus.objects.select_for_update().order_by("pk"))
            prepared = validate_rows(rows)
            save_prepared(prepared, created_by=request.user)
        request.session.pop("executor_import_rows", None)
        messages.success(request, f"Импортировано исполнителей: {len(prepared)}.")
        return redirect("executor_list")
    except CommandError as exc:
        request.session.pop("executor_import_rows", None)
        return render(
            request,
            "directory/executor_import.html",
            {"error": f"Импорт отменён: {exc}"},
        )

@login_required
def bulk_add_executors(request):
    if request.method != "POST":
        return redirect("executor_list")

    executor_ids = request.POST.getlist("executor_ids")
    project = get_object_or_404(Project, id=request.POST.get("project_id"))
    specialty = get_object_or_404(Specialty, id=request.POST.get("specialty_id"))
    new_status = get_object_or_404(SelectionStatus, name="Новый")
    return_url = request.POST.get("return_url") or "/executors/"
    if not return_url.startswith("/") or return_url.startswith("//"):
        return_url = "/executors/"

    if executor_ids:
        selected_executors = Executor.objects.filter(id__in=executor_ids)
        incompatible = selected_executors.exclude(
            executor_specialties__specialty=specialty
        ).distinct()
        compatible = selected_executors.exclude(id__in=incompatible.values_list("id", flat=True))
        if incompatible.exists():
            names = ", ".join(str(executor) for executor in incompatible[:5])
            suffix = " и другие" if incompatible.count() > 5 else ""
            messages.warning(
                request,
                f"Не добавлены исполнители без специальности {specialty.code}: "
                f"{names}{suffix}.",
            )

        if compatible.exists():
            with transaction.atomic():
                ProjectSpecialtyNeed.objects.get_or_create(
                    project=project,
                    specialty=specialty,
                    defaults={"created_by": request.user},
                )
                for executor_id in compatible.values_list("id", flat=True):
                    ProjectSelection.objects.get_or_create(
                        project=project,
                        specialty=specialty,
                        executor_id=executor_id,
                        defaults={"status": new_status},
                    )

    return redirect(return_url)

@login_required
def executor_list(request):
    optional_fields = [("city", "Город"), ("employment_type", "Тип занятости"),
                       ("birth_year", "Год рождения"), ("work_start_year", "Начал работать с"),
                       ("contact_source", "Источник контакта"), ("source_comment", "Комментарий к источнику"),
                       ("general_comment", "Общий комментарий"), ("status_comment", "Комментарий к статусу"),
                       ("software_comment", "Комментарий по ПО"), ("created_by", "Добавил")]
    specialty_ids = request.GET.getlist("specialty")
    software_ids = request.GET.getlist("software")
    search_query = request.GET.get("q", "").strip()
    sort = request.GET.get("sort", "name")

    # Здесь нужен именно статус исполнителя, а не статус подбора.
    executor_statuses = ExecutorStatus.objects.filter(is_active=True).order_by("name")
    if "status_filter" in request.GET:
        selected_status_ids = request.GET.getlist("status")
    else:
        selected_status_ids = [
            str(status_id)
            for status_id in executor_statuses.filter(
                name__in=("Активный", "Космонавт")
            ).values_list("id", flat=True)
        ]

    executors = (
        Executor.objects
        .select_related("status", "employment_type", "contact_source", "created_by")
        .prefetch_related(
            "executor_specialties__specialty",
            "software",
            "project_selections__project",
            "project_selections__specialty",
            "project_selections__status",
            "project_selections__project__status",
            "project_selections__negotiations__user",
            "comments",
        )
    )

    if specialty_ids:
        executors = executors.filter(
            executor_specialties__specialty_id__in=specialty_ids
        )

    if software_ids:
        executors = executors.filter(software__id__in=software_ids)

    if selected_status_ids:
        executors = executors.filter(status_id__in=selected_status_ids)

    if search_query:
        executors = executors.filter(
            models.Q(last_name__icontains=search_query)
            | models.Q(first_name__icontains=search_query)
            | models.Q(middle_name__icontains=search_query)
            | models.Q(phone__icontains=search_query)
            | models.Q(email__icontains=search_query)
            | models.Q(messenger__icontains=search_query)
            | models.Q(general_comment__icontains=search_query)
            | models.Q(status_comment__icontains=search_query)
            | models.Q(software_comment__icontains=search_query)
            | models.Q(city__icontains=search_query)
            | models.Q(contact_source__name__icontains=search_query)
            | models.Q(source_comment__icontains=search_query)
        )

    if sort == "status":
        executors = executors.order_by("status__name", "last_name", "first_name")
    elif sort == "created_desc":
        executors = executors.order_by("-created_at", "-id")
    elif sort == "created_asc":
        executors = executors.order_by("created_at", "id")
    else:
        executors = executors.order_by("last_name", "first_name", "middle_name")

    executors = executors.distinct()

    result_items = []

    for executor in executors:
        specialties_text = ", ".join(
            item.specialty.code
            for item in executor.executor_specialties.all()
        )
        software_text = ", ".join(item.name for item in executor.software.all())
        software_items = [item.name for item in executor.software.all()]

        selections = sorted(
            executor.project_selections.all(),
            key=lambda item: item.updated_at,
            reverse=True,
        )

        last_projects = []
        seen_project_ids = set()

        for selection in selections:
            if selection.project_id not in seen_project_ids:
                last_projects.append(
                    {
                        "project": selection.project.name,
                        "project_id": selection.project_id,
                        "specialty": selection.specialty.code,
                        "selection_status": selection.status.name if selection.status else "—",
                        "amount": format_money(selection.offer_amount),
            "status": selection.project.status.name if selection.project.status else "—",
                    }
                )
                seen_project_ids.add(selection.project_id)

            if len(last_projects) >= 3:
                break

        last_comment = executor.comments.all().order_by("-created_at").first()
        negotiations = [(note, selection.project) for selection in selections
                        for note in selection.negotiations.all()]
        latest = max(negotiations, key=lambda pair: (pair[0].created_at, pair[0].pk), default=None)

        result_items.append(
            {
                "executor": executor,
                "specialties": specialties_text,
                "software": software_text,
                "software_items": software_items,
                "last_projects": last_projects,
                "last_comment": last_comment,
                "last_negotiation": latest[0] if latest else None,
                "negotiation_project": latest[1] if latest else None,
                "optional_values": [getattr(executor, key) for key, label in optional_fields],
            }
        )

    context = {
        "items": result_items,
        "optional_fields": optional_fields,
        "specialties": Specialty.objects.filter(is_active=True).order_by("code"),
        "software_options": Software.objects.filter(is_active=True).order_by("name"),
        "executor_statuses": executor_statuses,
        "selected_specialty_ids": specialty_ids,
        "selected_software_ids": software_ids,
        "selected_status_ids": selected_status_ids,
        "search_query": search_query,
        "sort": sort,
        "projects": Project.objects.select_related("status").order_by("name"),
        "bulk_specialties": Specialty.objects.filter(is_active=True).order_by("code"),
        "total_executors_count": Executor.objects.count(),
        "active_executors_count": Executor.objects.filter(status__name="Активный").count(),
        "inactive_executors_count": Executor.objects.exclude(status__name="Активный").count(),
        "software_executors_count": Executor.objects.filter(
            software__isnull=False
        ).distinct().count(),
    }

    return render(request, "directory/executor_list.html", context)

@login_required
def executor_detail(request, executor_id):
    executor = get_object_or_404(
        Executor.objects
        .select_related("status", "employment_type", "created_by")
        .prefetch_related(
            "executor_specialties__specialty",
            "software",
            "comments__project",
            "comments__user",
            "project_selections__project",
            "project_selections__project__object_type",
            "project_selections__specialty",
            "project_selections__status",
            "project_selections__negotiations",
        ),
        id=executor_id,
    )

    specialties = [
        item.specialty
        for item in executor.executor_specialties.all()
    ]

    used_specialty_ids = [
        item.specialty_id
        for item in executor.executor_specialties.all()
    ]

    available_specialties = Specialty.objects.filter(
        is_active=True
    ).exclude(
        id__in=used_specialty_ids
    ).order_by("code")

    comments = executor.comments.all().order_by("-created_at")

    selections = sorted(
        executor.project_selections.all(),
        key=lambda item: item.updated_at,
        reverse=True,
    )

    selection_items = []

    for selection in selections:
        negotiations = selection.negotiations.all().order_by("-event_date", "-created_at")

        selection_items.append(
            {
                "selection": selection,
                "offer_amount": format_money(selection.offer_amount),
                "negotiations": negotiations,
            }
        )

    context = {
        "executor": executor,
        "specialties": specialties,
        "comments": comments,
        "selection_items": selection_items,
        "projects": Project.objects.order_by("name"),
        "executor_specialties": executor.executor_specialties.all(),
        "executor_statuses": ExecutorStatus.objects.filter(is_active=True),
        "available_specialties": available_specialties,
    }

    return render(request, "directory/executor_detail.html", context)


@login_required
@require_POST
def update_executor_status(request, executor_id):
    executor = get_object_or_404(Executor, pk=executor_id)
    if request.POST.get("field") == "status_comment":
        executor.status_comment = request.POST.get("value", "")
        executor.save(update_fields=["status_comment", "updated_at"])
        return JsonResponse({"saved": True})
    value = request.POST.get("status_id", "")
    status = ExecutorStatus.objects.filter(pk=value, is_active=True).first() if value.isdecimal() else None
    if status is None:
        return JsonResponse({"saved": False, "error": "Выберите доступный статус."}, status=400)
    if executor.status_id != status.pk:
        executor.status = status
        executor.status_changed_by = request.user
        executor.save(update_fields=["status", "status_changed_by", "updated_at"])
    author = executor.status_changed_by
    return JsonResponse({"saved": True, "status_id": status.pk,
                         "author": (author.get_full_name() or author.username) if author else "Не указан"})





@login_required
@require_GET
def executor_duplicates(request):
    surname = request.GET.get("last_name", "").strip()
    if not surname or len(surname) > 150:
        return JsonResponse({"matches": []})
    executors = Executor.objects.filter(last_name__iexact=surname).prefetch_related(
        "executor_specialties__specialty"
    ).order_by("first_name", "id")
    return JsonResponse({"matches": [{
        "id": item.pk, "name": str(item), "phone": item.phone or "—",
        "email": item.email or "—",
        "specialties": ", ".join(str(link.specialty) for link in item.executor_specialties.all()) or "—",
    } for item in executors]})


@login_required
def project_list(request):
    search_query = request.GET.get("q", "").strip()
    object_type_id = request.GET.get("object_type")

    projects = Project.objects.select_related(
        "object_type",
        "status",
        "chief_project_engineer",
        "created_by",
    ).order_by("name")

    selected_object_type = None

    if object_type_id:
        selected_object_type = ObjectType.objects.filter(id=object_type_id).first()

        if selected_object_type:
            projects = projects.filter(object_type=selected_object_type)

    if search_query:
        projects = projects.filter(
            models.Q(name__icontains=search_query)
            | models.Q(comment__icontains=search_query)
        )

    result_items = []

    for project in projects:
        selections = project.selections.all()

        total_offer_amount = sum(
            selection.offer_amount or 0
            for selection in selections
        )

        specialties = []
        seen_specialty_ids = set()

        for selection in selections:
            if selection.specialty_id not in seen_specialty_ids:
                specialties.append(selection.specialty.code)
                seen_specialty_ids.add(selection.specialty_id)

        result_items.append(
            {
                "project": project,
                "selections_count": len(selections),
                "total_offer_amount": format_money(total_offer_amount),
                "specialties": ", ".join(specialties),
            }
        )

    context = {
        "items": result_items,
        "search_query": search_query,
        "object_types": ObjectType.objects.filter(is_active=True).order_by("name"),
        "selected_object_type_id": object_type_id,
        "selected_object_type": selected_object_type,
        "total_projects_count": Project.objects.count(),
        "total_selections_count": ProjectSelection.objects.count(),
    }

    return render(request, "directory/project_list.html", context)

@login_required
def project_create(request):
    if request.method == "POST":
        form = ProjectForm(request.POST, user=request.user)

        if form.is_valid():
            project = form.save(commit=False)
            project.created_by = request.user

            if not project.chief_project_engineer:
                project.chief_project_engineer = request.user

            project.save()

            return redirect("project_detail", project_id=project.id)
    else:
        form = ProjectForm(user=request.user)

    context = {
        "form": form,
        "page_title": "Новый проект",
        "submit_text": "Создать проект",
        "back_url": "/projects/",
        "back_text": "← К списку проектов",
    }

    return render(request, "directory/project_form.html", context)

@login_required
def project_update(request, project_id):
    project = get_object_or_404(
        Project.objects.select_related(
            "object_type",
            "chief_project_engineer",
            "created_by",
        ),
        id=project_id,
    )

    if request.method == "POST":
        form = ProjectForm(
            request.POST,
            instance=project,
            user=request.user,
        )

        if form.is_valid():
            project = form.save()
            return redirect("project_detail", project_id=project.id)
    else:
        form = ProjectForm(
            instance=project,
            user=request.user,
        )

    context = {
        "form": form,
        "project": project,
        "page_title": f"Редактирование проекта: {project.name}",
        "submit_text": "Сохранить изменения",
        "back_url": f"/projects/{project.id}/",
        "back_text": "← К карточке проекта",
    }

    return render(request, "directory/project_form.html", context)

@login_required
def project_detail(request, project_id):
    project = get_object_or_404(
        Project.objects
        .select_related(
            "object_type",
            "status",
            "chief_project_engineer",
            "created_by",
        )
        .prefetch_related(
            "selections__specialty",
            "selections__executor",
            "selections__executor__status",
            "selections__status",
            "selections__negotiations",
            "selections__negotiations__user",
        ),
        id=project_id,
    )

    add_candidate_error = None

    selected_specialty_id = request.GET.get("specialty")
    selected_specialty = None
    candidate_items = []

    selection_specialty_id = request.GET.get("selection_specialty")
    selection_status_id = request.GET.get("selection_status")

    if request.method == "POST":
        action = request.POST.get("action")

        if action == "add_candidate":
            specialty_id = request.POST.get("specialty_id")
            executor_id = request.POST.get("executor_id")

            specialty = Specialty.objects.filter(id=specialty_id).first()
            executor = Executor.objects.filter(id=executor_id).first()
            new_status = SelectionStatus.objects.filter(name="Новый").first()

            if not new_status:
                add_candidate_error = (
                    "В справочнике статусов подбора нет статуса «Новый». "
                    "Добавьте его через админку."
                )
            elif not specialty or not executor:
                add_candidate_error = "Не выбрана специальность или исполнитель."
            elif not executor.executor_specialties.filter(specialty=specialty).exists():
                add_candidate_error = (
                    f"Исполнитель «{executor}» не имеет специальности {specialty.code} "
                    "и не может быть добавлен в этот раздел."
                )
            else:
                try:
                    ProjectSelection.objects.create(
                        project=project,
                        specialty=specialty,
                        executor=executor,
                        status=new_status,
                    )

                    need, created = ProjectSpecialtyNeed.objects.get_or_create(
                        project=project,
                        specialty=specialty,
                        defaults={
                            "created_by": request.user,
                        },
                    )

                    return redirect(
                        f"/projects/{project.id}/?specialty={specialty.id}#need-{need.id}"
                    )
                except IntegrityError:
                    add_candidate_error = (
                        "Этот исполнитель уже добавлен в проект по выбранной специальности."
                    )

            selected_specialty_id = specialty_id
            
        elif action == "add_specialty_need":
            specialty_id = request.POST.get("specialty_id")
            comment = request.POST.get("comment", "").strip()

            specialty = Specialty.objects.filter(id=specialty_id).first()

            if specialty:
                try:
                    ProjectSpecialtyNeed.objects.create(
                        project=project,
                        specialty=specialty,
                        comment=comment,
                        created_by=request.user,
                    )
                except IntegrityError:
                    pass

            return redirect("project_detail", project_id=project.id)
        
        else:
            return redirect("project_detail", project_id=project.id)
    
    if selected_specialty_id:
        selected_specialty = Specialty.objects.filter(id=selected_specialty_id).first()

        if selected_specialty:
            used_executor_ids = ProjectSelection.objects.filter(
                project=project,
                specialty=selected_specialty,
            ).values_list("executor_id", flat=True)

            candidates = (
                Executor.objects
                .filter(
                    executor_specialties__specialty=selected_specialty,
                    status__name="Активный",
                )
                .exclude(id__in=used_executor_ids)
                .select_related("status", "employment_type")
                .prefetch_related(
                    "executor_specialties__specialty",
                    "software",
                    "project_selections__project",
                    "project_selections__specialty",
                    "project_selections__status",
                    "comments",
                )
                .distinct()
                .order_by("last_name", "first_name", "middle_name")
            )

            for candidate in candidates:
                candidate_selections = [
                    selection
                    for selection in candidate.project_selections.all()
                    if selection.specialty_id == selected_specialty.id
                ]

                candidate_selections = sorted(
                    candidate_selections,
                    key=lambda item: item.updated_at,
                    reverse=True,
                )

                projects = []
                seen_project_ids = set()

                for selection in candidate_selections:
                    if selection.project_id not in seen_project_ids:
                        projects.append(
                            {
                                "name": selection.project.name,
                                "status": selection.project.status.name if selection.project.status else "—",
                                "specialty": selection.specialty.code,
                                "selection_status": selection.status.name if selection.status else "—",
                                "amount": format_money(selection.offer_amount),
                            }
                        )
                        seen_project_ids.add(selection.project_id)

                    if len(projects) >= 5:
                        break

                last_comment = candidate.comments.all().order_by("-created_at").first()

                candidate_items.append(
                    {
                        "executor": candidate,
                        "projects": projects,
                        "last_comment": last_comment,
                    }
                )

    selections = project.selections.all()

    if selection_specialty_id:
        selections = selections.filter(specialty_id=selection_specialty_id)

    if selection_status_id:
        selections = selections.filter(status_id=selection_status_id)

    selections = selections.order_by(
        "specialty__code",
        "offer_amount",
        "executor__last_name",
        "executor__first_name",
    )

    selection_items = []

    for selection in selections:
        negotiations = list(
            selection.negotiations.select_related("user").order_by(
                "-event_date",
                "-created_at",
            )
        )

        last_negotiation = negotiations[0] if negotiations else None
        old_negotiations = negotiations[1:] if len(negotiations) > 1 else []

        selection_items.append(
            {
                "selection": selection,
                "offer_amount": format_money(selection.offer_amount),
                "last_negotiation": last_negotiation,
                "old_negotiations": old_negotiations,
                "status_css_class": get_selection_status_css_class(
                    selection.status.name
                ),
            }
        )
        
    needs = (
    project.specialty_needs
    .select_related("specialty", "created_by")
    .order_by("specialty__code")
    )

    all_project_selections = list(
        project.selections
        .select_related("specialty", "executor", "executor__status", "status")
        .prefetch_related("negotiations", "attachments__uploaded_by")
        .all()
    )

    need_items = []
    grouped_need_items = []

    project_min_total = 0
    project_max_total = 0
    project_sections_with_amount = 0

    needs_closed_count = 0
    needs_in_progress_count = 0
    needs_empty_count = 0

    selection_status_counters = {
        "Новый": 0,
        "Рассматривает": 0,
        "Готов": 0,
        "Утвержден": 0,
        "Отказ": 0,
        "Другие": 0,
    }

    for selection in all_project_selections:
        status_name = selection.status.name if selection.status else ""

        if status_name in selection_status_counters:
            selection_status_counters[status_name] += 1
        else:
            selection_status_counters["Другие"] += 1

    for need in needs:
        need_selections = [
            selection
            for selection in all_project_selections
            if selection.specialty_id == need.specialty_id
        ]

        offer_amounts = [
            selection.offer_amount
            for selection in need_selections
            if selection.offer_amount is not None
        ]

        min_offer_amount = min(offer_amounts) if offer_amounts else None
        max_offer_amount = max(offer_amounts) if offer_amounts else None

        if min_offer_amount is not None:
            project_min_total += min_offer_amount
            project_max_total += max_offer_amount
            project_sections_with_amount += 1

        ready_count = sum(
            1
            for selection in need_selections
            if selection.status and selection.status.name == "Утвержден"
        )
        ready_selection_count = sum(
            1
            for selection in need_selections
            if selection.status and selection.status.name in {"Готов", "Утвержден"}
        )

        need_status = get_need_status(need_selections)

        if need_status["label"] == "Закрыт":
            needs_closed_count += 1
        elif need_status["label"] == "В работе":
            needs_in_progress_count += 1
        else:
            needs_empty_count += 1

        selection_rows = []

        for selection in sorted(
            need_selections,
            key=lambda item: (
                item.status.name if item.status else "",
                item.executor.last_name,
                item.executor.first_name,
            ),
        ):
            negotiations = list(
                selection.negotiations.select_related("user").order_by(
                    "-event_date",
                    "-created_at",
                )
            )

            last_negotiation = negotiations[0] if negotiations else None
            old_negotiations = negotiations[1:] if len(negotiations) > 1 else []

            selection_rows.append(
                {
                    "selection": selection,
                    "offer_amount": format_money(selection.offer_amount),
                    "last_negotiation": last_negotiation,
                    "old_negotiations": old_negotiations,
                    "status_css_class": get_selection_status_css_class(
                        selection.status.name
                    ),
                }
            )
            
        is_candidate_section_active = (
            selected_specialty_id
            and str(selected_specialty_id) == str(need.specialty_id)
        )

        need_item = {
            "need": need,
            "selections": need_selections,
            "selection_rows": selection_rows,
            "selections_count": len(need_selections),
            "ready_count": ready_count,
            "ready_selection_count": ready_selection_count,
            "min_offer_amount": format_money(min_offer_amount),
            "max_offer_amount": format_money(max_offer_amount),
            "status_label": need_status["label"],
            "status_css_class": need_status["css_class"],
            "is_candidate_section_active": is_candidate_section_active,
            "candidate_items": candidate_items if is_candidate_section_active else [],
        }

        need_items.append(need_item)
        grouped_need_items.append(need_item)

    selection_status_counters = {
        "Новый": 0,
        "Рассматривает": 0,
        "Готов": 0,
        "Утвержден": 0,
        "Отказ": 0,
        "Другие": 0,
    }
    for item in need_items:
        for selection in item["selections"]:
            status_name = selection.status.name if selection.status else ""
            if status_name in selection_status_counters:
                selection_status_counters[status_name] += 1
            else:
                selection_status_counters["Другие"] += 1

    used_need_specialty_ids = [
        item["need"].specialty_id
        for item in need_items
    ]

    available_need_specialties = (
        Specialty.objects
        .filter(is_active=True)
        .exclude(id__in=used_need_specialty_ids)
        .order_by("code")
    )

    used_need_specialty_ids = [
        item["need"].specialty_id
        for item in need_items
    ]

    available_need_specialties = (
        Specialty.objects
        .filter(is_active=True)
        .exclude(id__in=used_need_specialty_ids)
        .order_by("code")
    )    

    context = {
        "project": project,
        "selection_items": selection_items,
        "selections_count": sum(item["selections_count"] for item in need_items),
        "project_min_total": format_money(project_min_total),
        "project_max_total": format_money(project_max_total),
        "project_sections_with_amount": project_sections_with_amount,
        
        "need_items": need_items,
        "available_need_specialties": available_need_specialties,
        "needs_count": len(need_items),
        "closed_needs_count": sum(
            1 for item in need_items if item["ready_count"] > 0
        ),
        
        "add_candidate_error": add_candidate_error,
        "specialties": Specialty.objects.filter(is_active=True).order_by("code"),
        "selected_specialty_id": selected_specialty_id,
        "selected_specialty": selected_specialty,
        "candidate_items": candidate_items,
        "selection_statuses": SelectionStatus.objects.filter(is_active=True).order_by(
            models.Case(
                models.When(name="Новый", then=0),
                models.When(name="Рассматривает", then=1),
                models.When(name="Отказ", then=2),
                models.When(name="Готов", then=3),
                models.When(name="Утвержден", then=4),
                default=99,
                output_field=models.IntegerField(),
            ),
            "name",
        ),
        "selection_specialty_id": selection_specialty_id,
        "selection_status_id": selection_status_id,
        
        "grouped_need_items": grouped_need_items,

        "project_min_total": format_money(project_min_total),
        "project_max_total": format_money(project_max_total),
        "project_sections_with_amount": project_sections_with_amount,

        "needs_count": len(need_items),
        "closed_needs_count": needs_closed_count,
        "needs_in_progress_count": needs_in_progress_count,
        "needs_empty_count": needs_empty_count,

        "selection_status_counters": selection_status_counters,
        "project_statuses": ProjectStatus.objects.filter(is_active=True).order_by(
            models.Case(
                models.When(name="Черновик", then=0),
                models.When(name="В работе", then=1),
                models.When(name="Завершен", then=2),
                default=99,
                output_field=models.IntegerField(),
            ),
            "name",
        ),
    }

    return render(request, "directory/project_detail.html", context)


@login_required
def update_project_status(request, project_id):
    project = get_object_or_404(Project, id=project_id)

    if request.method == "POST":
        status = ProjectStatus.objects.filter(
            id=request.POST.get("status_id"),
            is_active=True,
        ).first()
        if status:
            project.status = status
            project.save(update_fields=["status", "updated_at"])

    return redirect("project_detail", project_id=project.id)

@login_required
def update_project_selection(request, selection_id):
    selection = get_object_or_404(
        ProjectSelection.objects.select_related("project"),
        id=selection_id,
    )

    project_id = selection.project.id

    if request.method != "POST":
        return redirect(f"/projects/{project_id}/#selection-{selection.id}")

    if request.POST.get("autosave") == "1":
        field = request.POST.get("field")
        value = request.POST.get("value", "")

        if field == "status_id":
            status = SelectionStatus.objects.filter(
                id=value,
                is_active=True,
            ).first()

            if not status:
                return JsonResponse(
                    {"saved": False, "error": "Выбран недоступный статус."},
                    status=400,
                )

            selection.status = status
            selection.save(update_fields=["status", "updated_at"])

            section_selections = list(ProjectSelection.objects.filter(
                project_id=project_id, specialty_id=selection.specialty_id,
            ).select_related("status"))
            section_status = get_need_status(section_selections)

            return JsonResponse(
                {
                    "saved": True,
                    "field": field,
                    "value": str(status.id),
                    "status_css_class": get_selection_status_css_class(status.name),
                    "status_name": status.name,
                    "section_status": section_status,
                    "approved_count": sum(
                        bool(item.status and item.status.name == "Утвержден")
                        for item in section_selections
                    ),
                    "ready_count": sum(
                        bool(
                            item.status
                            and item.status.name in {"Готов", "Утвержден"}
                        )
                        for item in section_selections
                    ),
                }
            )

        if field == "offer_amount":
            normalized_amount = value.replace(" ", "").replace("\u00a0", "")

            if normalized_amount:
                if not normalized_amount.isdecimal():
                    return JsonResponse(
                        {"saved": False, "error": "КП должно быть целым числом."},
                        status=400,
                    )

                amount = int(normalized_amount)
                if amount > 2_147_483_647:
                    return JsonResponse(
                        {"saved": False, "error": "Указана слишком большая сумма КП."},
                        status=400,
                    )
                selection.offer_amount = amount
            else:
                selection.offer_amount = None

            selection.save(update_fields=["offer_amount", "updated_at"])

            return JsonResponse(
                {
                    "saved": True,
                    "field": field,
                    "value": (
                        ""
                        if selection.offer_amount is None
                        else str(selection.offer_amount)
                    ),
                    "display_value": format_money(selection.offer_amount) or "",
                }
            )

        if field == "comment":
            selection.comment = value.strip()
            selection.save(update_fields=["comment", "updated_at"])

            return JsonResponse(
                {
                    "saved": True,
                    "field": field,
                    "value": selection.comment,
                }
            )

        return JsonResponse(
            {"saved": False, "error": "Это поле нельзя сохранить автоматически."},
            status=400,
        )

    status_id = request.POST.get("status_id")
    offer_amount_raw = request.POST.get("offer_amount", "").strip()
    comment = request.POST.get("comment", "").strip()

    if status_id:
        status = SelectionStatus.objects.filter(id=status_id).first()
        if status:
            selection.status = status

    if offer_amount_raw:
        normalized_amount = (
            offer_amount_raw
            .replace(" ", "")
            .replace("\u00a0", "")
            .replace(",", ".")
        )

        try:
            selection.offer_amount = int(float(normalized_amount))
        except ValueError:
            pass
    else:
        selection.offer_amount = None

    selection.comment = comment
    selection.save()

    return redirect(f"/projects/{project_id}/#selection-{selection.id}")

@login_required
def add_selection_negotiation(request, selection_id):
    selection = get_object_or_404(
        ProjectSelection.objects.select_related("project"),
        id=selection_id,
    )

    if request.method != "POST":
        return redirect("project_detail", project_id=selection.project.id)

    comment = request.POST.get("comment", "").strip()

    if comment:
        SelectionNegotiation.objects.create(
            selection=selection,
            event_date=timezone.localdate(),
            comment=comment,
            user=request.user,
        )

    specialty_id = request.POST.get("specialty")
    suffix = f"?specialty={specialty_id}" if specialty_id else ""
    return redirect(f"/projects/{selection.project.id}/{suffix}#selection-{selection.id}")

@login_required
def delete_project_selection(request, selection_id):
    selection = get_object_or_404(
        ProjectSelection.objects.select_related("project"),
        id=selection_id,
    )

    project_id = selection.project.id

    if request.method == "POST":
        selection.delete()

    return redirect("project_detail", project_id=project_id)

@login_required
def add_executor_comment(request, executor_id):
    executor = get_object_or_404(Executor, id=executor_id)

    if request.method != "POST":
        return redirect("executor_detail", executor_id=executor.id)

    project_id = request.POST.get("project_id")
    comment_text = request.POST.get("comment", "").strip()

    project = None

    if project_id:
        project = Project.objects.filter(id=project_id).first()

    if comment_text:
        ExecutorComment.objects.create(
            executor=executor,
            project=project,
            user=request.user,
            comment=comment_text,
        )

    return redirect("executor_detail", executor_id=executor.id)

def update_executor_specialties(executor, specialties):
    specialties = specialties or []

    ExecutorSpecialty.objects.filter(executor=executor).exclude(
        specialty__in=specialties
    ).delete()

    for specialty in specialties:
        ExecutorSpecialty.objects.get_or_create(
            executor=executor,
            specialty=specialty,
        )

@login_required
def executor_create(request):
    if request.method == "POST":
        form = ExecutorForm(request.POST)

        if form.is_valid():
            form.instance.created_by = request.user
            executor = form.save()
            specialties = form.cleaned_data.get("specialties")
            update_executor_specialties(executor, specialties)

            return redirect("executor_detail", executor_id=executor.id)
    else:
        form = ExecutorForm()

    context = {
        "form": form,
        "page_title": "Новый исполнитель",
        "submit_text": "Создать исполнителя",
        "back_url": "/executors/",
        "back_text": "← К списку исполнителей",
    }

    return render(request, "directory/executor_form.html", context)

@login_required
def add_executor_specialty(request, executor_id):
    executor = get_object_or_404(Executor, id=executor_id)

    if request.method != "POST":
        return redirect("executor_detail", executor_id=executor.id)

    specialty_id = request.POST.get("specialty")
    specialty = Specialty.objects.filter(id=specialty_id, is_active=True).first()

    if specialty:
        ExecutorSpecialty.objects.get_or_create(
            executor=executor,
            specialty=specialty,
        )

    return redirect("executor_detail", executor_id=executor.id)


@login_required
def delete_executor_specialty(request, executor_specialty_id):
    executor_specialty = get_object_or_404(
        ExecutorSpecialty.objects.select_related("executor"),
        id=executor_specialty_id,
    )

    executor_id = executor_specialty.executor.id

    if request.method == "POST":
        executor_specialty.delete()

    return redirect("executor_detail", executor_id=executor_id)

@login_required
def executor_update(request, executor_id):
    executor = get_object_or_404(
        Executor.objects
        .select_related("status", "employment_type")
        .prefetch_related("executor_specialties__specialty"),
        id=executor_id,
    )

    selected_specialty_ids = list(
        executor.executor_specialties.values_list(
            "specialty_id",
            flat=True,
        )
    )

    if request.method == "POST":
        form = ExecutorForm(request.POST, instance=executor)

        if form.is_valid():
            if "status" in form.changed_data:
                form.instance.status_changed_by = request.user
            executor = form.save()
            specialties = form.cleaned_data.get("specialties")
            update_executor_specialties(executor, specialties)

            return redirect("executor_detail", executor_id=executor.id)
    else:
        form = ExecutorForm(
            instance=executor,
            initial={
                "specialties": selected_specialty_ids,
            },
        )

    context = {
        "form": form,
        "executor": executor,
        "page_title": f"Редактирование: {executor}",
        "submit_text": "Сохранить изменения",
        "back_url": f"/executors/{executor.id}/",
        "back_text": "← К карточке исполнителя",
    }

    return render(request, "directory/executor_form.html", context)


@login_required
@require_http_methods(["GET", "POST"])
def executor_delete(request, executor_id):
    executor = get_object_or_404(Executor, id=executor_id)

    if request.method == "POST":
        executor.delete()
        messages.success(request, "Исполнитель удалён.")
        return redirect("executor_list")

    return render(
        request,
        "directory/executor_delete_confirm.html",
        {"executor": executor},
    )

@login_required
@require_POST
def delete_selection_negotiation(request, negotiation_id):
    negotiation = get_object_or_404(
        SelectionNegotiation.objects.select_related("selection__project"),
        id=negotiation_id,
    )

    if not (request.user.pk == negotiation.user_id or request.user.is_staff or request.user.is_superuser):
        return HttpResponseForbidden("Изменять запись переговоров может только её автор или администратор.")

    project_id = negotiation.selection.project.id

    if request.method == "POST":
        negotiation.delete()

    return redirect("project_detail", project_id=project_id)


@login_required
@require_POST
def edit_selection_negotiation(request, negotiation_id):
    negotiation = get_object_or_404(SelectionNegotiation.objects.select_related("selection"), pk=negotiation_id)
    if not (request.user.pk == negotiation.user_id or request.user.is_staff or request.user.is_superuser):
        return HttpResponseForbidden("Изменять запись переговоров может только её автор или администратор.")
    comment = request.POST.get("comment", "").strip()
    if not comment:
        return HttpResponseBadRequest("Введите текст переговоров.")
    negotiation.comment = comment
    negotiation.save(update_fields=["comment"])
    return redirect("project_detail", project_id=negotiation.selection.project_id)

@login_required
def delete_project_specialty_need(request, need_id):
    need = get_object_or_404(
        ProjectSpecialtyNeed.objects.select_related("project"),
        id=need_id,
    )

    project_id = need.project.id

    if request.method == "POST":
        with transaction.atomic():
            ProjectSelection.objects.filter(
                project_id=need.project_id,
                specialty_id=need.specialty_id,
            ).delete()
            need.delete()

    return redirect("project_detail", project_id=project_id)


ATTACHMENT_EXTENSIONS = {".pdf", ".doc", ".docx", ".xls", ".xlsx", ".jpg", ".jpeg", ".png", ".zip"}
ATTACHMENT_MAX_SIZE = 20 * 1024 * 1024


def attachment_selection_url(selection):
    return f"/projects/{selection.project_id}/?specialty={selection.specialty_id}#selection-{selection.pk}"


@login_required
@require_POST
def upload_selection_attachments(request, selection_id):
    selection = get_object_or_404(ProjectSelection, pk=selection_id)
    uploads = request.FILES.getlist("files")
    errors = []
    if not uploads:
        errors.append("Выберите файлы для загрузки.")
    for upload in uploads:
        if Path(upload.name).suffix.lower() not in ATTACHMENT_EXTENSIONS:
            errors.append(f"{upload.name}: неподдерживаемый формат файла.")
        if upload.size > ATTACHMENT_MAX_SIZE:
            errors.append(f"{upload.name}: размер файла не должен превышать 20 МБ.")
        if not upload.size:
            errors.append(f"{upload.name}: файл пустой.")
    if errors:
        for error in errors:
            messages.error(request, error)
        return redirect(attachment_selection_url(selection))

    saved_files = []
    try:
        with transaction.atomic():
            for upload in uploads:
                attachment = SelectionAttachment(
                    selection=selection, original_name=upload.name,
                    size=upload.size, uploaded_by=request.user,
                )
                attachment.file.save(upload.name, upload, save=False)
                saved_files.append((attachment.file.storage, attachment.file.name))
                attachment.save()
    except Exception:
        # Storage writes are not rolled back by the database transaction.
        for storage, name in saved_files:
            storage.delete(name)
        raise
    messages.success(request, f"Загружено файлов: {len(uploads)}.")
    return redirect(attachment_selection_url(selection))


@login_required
@require_GET
def download_selection_attachment(request, attachment_id):
    attachment = get_object_or_404(SelectionAttachment, pk=attachment_id)
    try:
        stream = attachment.file.open("rb")
    except FileNotFoundError:
        raise Http404("Файл не найден в хранилище.") from None
    return FileResponse(
        stream, as_attachment=True, filename=attachment.original_name,
        content_type="application/octet-stream",
    )


@login_required
@require_POST
def delete_selection_attachment(request, attachment_id):
    attachment = get_object_or_404(SelectionAttachment.objects.select_related("selection"), pk=attachment_id)
    target = attachment_selection_url(attachment.selection)
    attachment.delete()
    messages.success(request, "Файл удалён.")
    return redirect(target)
